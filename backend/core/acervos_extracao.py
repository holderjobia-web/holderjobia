"""Extração de texto de arquivos do módulo Acervos, best-effort.

GOVERNANÇA/ESCOPO: só extrai o que dá pra ler como texto (PDF, XLSX, CSV,
TXT/MD) usando as bibliotecas já presentes no projeto (pdfplumber, openpyxl).
Fotos, plantas em CAD, Word/PowerPoint e outros formatos binários sem parser
disponível ficam SEM TEXTO — isso NÃO é erro: o arquivo continua guardado e
visível no módulo, só não entra na base de conhecimento (RAG) do agente.
"""

import csv
import io
import logging
from typing import Optional

logger = logging.getLogger(__name__)

_EXTENSOES_TEXTO_PLANO = ("txt", "md")


def extensao_de(nome_arquivo: str) -> str:
    """Extensão em minúscula, sem ponto (ex.: 'pdf'); vazio se não houver."""
    return nome_arquivo.rsplit(".", 1)[-1].lower() if "." in nome_arquivo else ""


def extrair_texto(nome_arquivo: str, conteudo: bytes) -> Optional[str]:
    """Texto extraído do arquivo, ou None se a extensão não tem parser disponível."""
    extensao = extensao_de(nome_arquivo)
    try:
        if extensao == "pdf":
            from core.dre_parser import extrair_texto_pdf
            return extrair_texto_pdf(conteudo).strip() or None

        if extensao == "xlsx":
            return _extrair_texto_xlsx(conteudo) or None

        if extensao == "csv":
            return _extrair_texto_csv(conteudo) or None

        if extensao in _EXTENSOES_TEXTO_PLANO:
            return conteudo.decode("utf-8", errors="ignore").strip() or None
    except Exception:
        logger.exception("Falha ao extrair texto (arquivo=%s extensao=%s)", nome_arquivo, extensao)
        return None

    return None  # extensão sem parser disponível (imagem, dwg, docx, etc.)


def _extrair_texto_xlsx(conteudo: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    partes: list[str] = []
    for ws in wb.worksheets:
        for linha in ws.iter_rows(values_only=True):
            celulas = [str(c) for c in linha if c is not None]
            if celulas:
                partes.append(" | ".join(celulas))
    wb.close()
    return "\n".join(partes)


def _ler_csv(conteudo: bytes) -> list[list[str]]:
    texto = conteudo.decode("utf-8-sig", errors="ignore")
    return [linha for linha in csv.reader(io.StringIO(texto)) if any(linha)]


def _extrair_texto_csv(conteudo: bytes) -> str:
    return "\n".join(" | ".join(linha) for linha in _ler_csv(conteudo))


def previa_csv(conteudo: bytes, limite: int = 25) -> dict:
    """Primeiras linhas do CSV, mesmo formato de prévia usado no módulo Vendas."""
    linhas = _ler_csv(conteudo)
    colunas = linhas[0] if linhas else []
    amostra = linhas[1:1 + limite]
    return {"colunas": colunas, "linhas": amostra, "total_linhas": max(len(linhas) - 1, 0)}
