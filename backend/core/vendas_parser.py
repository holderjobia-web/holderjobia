"""Parser de planilhas de Vendas (.xlsx) — export mensal da clínica.

Mesma governança do parser de DRE (core/dre_parser.py):
- Extração por RÓTULO: localiza a linha de cabeçalho real (1ª coluna ==
  "Paciente"), não assume posição fixa — sobrevive à linha de endereço da
  clínica que vem antes dos dados no export.
- Cada linha da planilha é 1 venda/parcela recebida. O MÊS de referência é
  DERIVADO da coluna "Pagamento" (data por linha) — nunca do campo informado
  manualmente no upload, que é só um rótulo/metadado do envio.
- A CATEGORIA do procedimento é DERIVADA da coluna "Dentista", que traz o
  procedimento como sufixo do nome (ex.: "DRA DENISE ALVES SOUZA ORTO").
  Um único arquivo mensal pode conter todas as categorias misturadas.
- FAIL-SAFE: linha sem paciente/data/valor recebido válidos é ignorada e
  contada em `linhas_invalidas`; nunca fabrica valor nem derruba o arquivo
  inteiro por causa de 1 linha ruim. Isso também descarta naturalmente o
  rodapé de totais do export ("TOTAL GERAL", "BRASILCARD", "PIX"...), que vem
  sem data de pagamento.
- Procedimento não reconhecido vira 'nao_identificado' (sinalizado), nunca é
  descartado nem chutado para outra categoria.
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Optional

from openpyxl import load_workbook

_COLUNA_HEADER_ESPERADA = "Paciente"

CATEGORIA_NAO_IDENTIFICADA = "nao_identificado"

# Token(s) que aparecem no nome do dentista -> categoria do procedimento.
# Ordem importa: a primeira regra que casar vence (mais específica primeiro).
_REGRAS_CATEGORIA: tuple[tuple[tuple[str, ...], str], ...] = (
    (("IMPLANTE",), "implante"),
    (("IMPLANTODONTIA",), "implante"),
    (("ENDO",), "endodontia"),
    (("ENDODONTIA",), "endodontia"),
    (("RADIOLOGIA",), "radiologia"),
    (("ORTO",), "ortodontia"),
    (("ORTODONTIA",), "ortodontia"),
    (("CLINICO", "GERAL"), "clinico_geral"),
    (("C", "GERAL"), "clinico_geral"),
)

LABEL_CATEGORIA = {
    "ortodontia": "Ortodontia",
    "clinico_geral": "Clínico geral",
    "implante": "Implante",
    "endodontia": "Endodontia",
    "radiologia": "Radiologia",
    CATEGORIA_NAO_IDENTIFICADA: "Não identificado",
}


@dataclass
class MesAgregado:
    mes_referencia: date
    categoria: str
    quantidade_vendas: int
    valor_original: Optional[Decimal]
    valor_desconto: Optional[Decimal]
    valor_recebido: Decimal
    por_forma_pagamento: dict = field(default_factory=dict)


@dataclass
class ResultadoParseVendas:
    ok: bool
    meses: list[MesAgregado] = field(default_factory=list)
    linhas_validas: int = 0
    linhas_invalidas: int = 0
    nao_identificados: list[str] = field(default_factory=list)
    motivo: Optional[str] = None


def _normalizar(texto: str) -> str:
    """Uppercase sem acento — base da comparação de tokens."""
    sem_acento = "".join(
        c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)
    )
    return sem_acento.upper().strip()


def classificar_categoria(dentista) -> str:
    """Deriva a categoria do procedimento a partir do nome do dentista.

    Casa por TOKEN (palavra inteira), não por substring — evita falso positivo
    tipo 'ENDO' dentro de um sobrenome. Sem match = 'nao_identificado'.
    """
    if not dentista:
        return CATEGORIA_NAO_IDENTIFICADA
    tokens = _normalizar(str(dentista)).split()
    if not tokens:
        return CATEGORIA_NAO_IDENTIFICADA

    for termos, categoria in _REGRAS_CATEGORIA:
        n = len(termos)
        if n == 1:
            if termos[0] in tokens:
                return categoria
        elif any(tuple(tokens[i:i + n]) == termos for i in range(len(tokens) - n + 1)):
            return categoria
    return CATEGORIA_NAO_IDENTIFICADA


def _parse_valor_brl(bruto) -> Optional[Decimal]:
    """'R$ 1.234,56' -> Decimal('1234.56'). None/vazio -> None."""
    if bruto is None:
        return None
    s = str(bruto).strip()
    if not s:
        return None
    s = s.replace("R$", "").strip().replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _parse_data_br(bruto) -> Optional[date]:
    """'31/08/2026' -> date(2026,8,31). Aceita datetime já convertido pelo openpyxl."""
    if bruto is None:
        return None
    if isinstance(bruto, datetime):
        return bruto.date()
    if isinstance(bruto, date):
        return bruto
    s = str(bruto).strip()
    try:
        dia, mes, ano = s.split("/")
        return date(int(ano), int(mes), int(dia))
    except (ValueError, AttributeError):
        return None


def _localizar_header(linhas: list[tuple]) -> Optional[int]:
    for i, linha in enumerate(linhas):
        if not linha:
            continue
        primeira = linha[0]
        if isinstance(primeira, str) and primeira.strip() == _COLUNA_HEADER_ESPERADA:
            return i
    return None


def previa_planilha(conteudo: bytes, limite: int = 25) -> dict:
    """Primeiras linhas da planilha, para conferir visualmente qual arquivo é.

    O .xlsx não abre direto no navegador, então a prévia é montada aqui. A 1ª
    linha do export traz o nome/endereço da unidade — é o que identifica o
    arquivo de relance.
    """
    wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    total_linhas = len(linhas)
    wb.close()

    idx_header = _localizar_header(linhas)
    identificacao = None
    if linhas and linhas[0] and isinstance(linhas[0][0], str):
        identificacao = linhas[0][0].strip()

    colunas: list[str] = []
    if idx_header is not None:
        colunas = [str(c).strip() if c is not None else "" for c in linhas[idx_header]]

    inicio = (idx_header + 1) if idx_header is not None else 0
    amostra = [
        ["" if c is None else str(c) for c in linha]
        for linha in linhas[inicio:inicio + limite]
    ]

    return {
        "identificacao": identificacao,
        "colunas": colunas,
        "linhas": amostra,
        "total_linhas": total_linhas,
    }


def parsear_planilha(conteudo: bytes) -> ResultadoParseVendas:
    """Parseia o .xlsx de vendas e agrega por (mês, categoria), ambos derivados
    linha a linha (coluna "Pagamento" e coluna "Dentista")."""
    try:
        wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        linhas = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return ResultadoParseVendas(
            ok=False, motivo="Não foi possível ler o arquivo (.xlsx inválido ou corrompido)."
        )

    idx_header = _localizar_header(linhas)
    if idx_header is None:
        return ResultadoParseVendas(
            ok=False,
            motivo=f"Cabeçalho não encontrado (esperava uma linha com '{_COLUNA_HEADER_ESPERADA}' na 1ª coluna).",
        )

    grupos: dict[tuple[date, str], list[dict]] = {}
    nao_identificados: set[str] = set()
    validas = 0
    invalidas = 0

    for linha in linhas[idx_header + 1:]:
        if not linha or all(v is None for v in linha):
            continue

        paciente = linha[0] if len(linha) > 0 else None
        data_pagamento = _parse_data_br(linha[3]) if len(linha) > 3 else None
        valor_recebido = _parse_valor_brl(linha[5]) if len(linha) > 5 else None

        if not paciente or data_pagamento is None or valor_recebido is None:
            invalidas += 1
            continue

        valor_original = _parse_valor_brl(linha[2]) if len(linha) > 2 else None
        valor_desconto = _parse_valor_brl(linha[4]) if len(linha) > 4 else None
        tipo_pagamento_bruto = linha[6] if len(linha) > 6 else None
        tipo_pagamento = (
            tipo_pagamento_bruto.strip() if isinstance(tipo_pagamento_bruto, str) and tipo_pagamento_bruto.strip()
            else "Não informado"
        )
        dentista = linha[8] if len(linha) > 8 else None
        categoria = classificar_categoria(dentista)
        if categoria == CATEGORIA_NAO_IDENTIFICADA and dentista:
            nao_identificados.add(str(dentista).strip())

        mes_ref = data_pagamento.replace(day=1)
        grupos.setdefault((mes_ref, categoria), []).append({
            "valor_original": valor_original,
            "valor_desconto": valor_desconto,
            "valor_recebido": valor_recebido,
            "tipo_pagamento": tipo_pagamento,
        })
        validas += 1

    if validas == 0:
        return ResultadoParseVendas(
            ok=False,
            motivo="Nenhuma linha válida encontrada (verifique paciente/data/valor recebido).",
            linhas_invalidas=invalidas,
        )

    meses: list[MesAgregado] = []
    for chave in sorted(grupos):
        mes_ref, categoria = chave
        vendas_mes = grupos[chave]
        soma_original = Decimal("0")
        soma_desconto = Decimal("0")
        soma_recebido = Decimal("0")
        tem_original = False
        tem_desconto = False
        por_forma: dict[str, Decimal] = {}

        for v in vendas_mes:
            if v["valor_original"] is not None:
                soma_original += v["valor_original"]
                tem_original = True
            if v["valor_desconto"] is not None:
                soma_desconto += v["valor_desconto"]
                tem_desconto = True
            soma_recebido += v["valor_recebido"]
            forma = v["tipo_pagamento"]
            por_forma[forma] = por_forma.get(forma, Decimal("0")) + v["valor_recebido"]

        meses.append(MesAgregado(
            mes_referencia=mes_ref,
            categoria=categoria,
            quantidade_vendas=len(vendas_mes),
            valor_original=soma_original if tem_original else None,
            valor_desconto=soma_desconto if tem_desconto else None,
            valor_recebido=soma_recebido,
            por_forma_pagamento=por_forma,
        ))

    return ResultadoParseVendas(
        ok=True,
        meses=meses,
        linhas_validas=validas,
        linhas_invalidas=invalidas,
        nao_identificados=sorted(nao_identificados),
    )
